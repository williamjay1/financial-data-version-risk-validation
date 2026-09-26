"""Render the completed Markdown manuscript into one anonymous Word manuscript.

Run with the bundled document Python after the artifact-operation marker.
Pandoc handles citations and editable Office equations; python-docx sets layout.
The generated Word file must subsequently pass render_docx.py and page review.
"""
import argparse
import json
import re
import shutil
import subprocess
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
MAN=ROOT/'manuscript'

def set_element(parent,tag,attrs):
    el=parent.find(qn(tag))
    if el is None:el=OxmlElement(tag);parent.append(el)
    for key,value in attrs.items():el.set(qn(key),str(value))
    return el

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',type=Path,default=MAN/'manuscript.md')
    parser.add_argument('--output',type=Path,default=MAN/'Auditing_Financial_Data_Versions_Revised.docx')
    args=parser.parse_args()
    if args.output.resolve().drive.lower()!='d:':raise ValueError('Document outputs belong on D:')
    source=args.source.read_text(encoding='utf-8')
    if re.search(r'\b(TODO|TBD|INSERT RESULT|PLACEHOLDER)\b',source,re.I):raise ValueError('Unresolved manuscript placeholders')
    exe=shutil.which('pandoc')
    if not exe:raise FileNotFoundError('Pandoc executable unavailable')
    (ROOT/'temp').mkdir(exist_ok=True)
    intermediate=ROOT/'temp/manuscript_cited.docx'
    command=[exe,str(args.source),'--standalone','--citeproc','--bibliography='+str(MAN/'references.bib'),
             '--resource-path='+str(ROOT)+';'+str(MAN),'-M','lang=en-US','-o',str(intermediate)]
    run=subprocess.run(command,capture_output=True,text=True,encoding='utf-8',cwd=ROOT)
    (ROOT/'temp/pandoc_build.log').write_text(run.stdout+'\n'+run.stderr,encoding='utf-8')
    if run.returncode:raise RuntimeError(run.stderr)
    if 'Could not convert TeX math' in run.stderr or 'Citeproc: citation' in run.stderr:raise RuntimeError('Native equation conversion failed: '+run.stderr)
    doc=Document(intermediate)
    for section in doc.sections:
        section.page_width=Inches(8.5);section.page_height=Inches(11)
        section.top_margin=Inches(.82);section.bottom_margin=Inches(.82)
        section.left_margin=Inches(.85);section.right_margin=Inches(.85)
        section.header_distance=Inches(.32);section.footer_distance=Inches(.32)
        footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        footer.text=''
        r=footer.add_run();r.font.name='Times New Roman';r.font.size=Pt(9)
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
    for style in doc.styles:
        if style.type in (1,2):
            style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
            rf=style.element.find(qn('w:rPr'))
            if rf is not None:
                color=rf.find(qn('w:color'))
                if color is not None:
                    for a in ['w:themeColor','w:themeTint','w:themeShade']:color.attrib.pop(qn(a),None)
    for name in ['Normal','Body Text','First Paragraph']:
        if name in doc.styles:
            st=doc.styles[name];st.font.size=Pt(11)
            st.paragraph_format.line_spacing=1.14;st.paragraph_format.space_after=Pt(6)
            st.paragraph_format.widow_control=True
    for name,size in [('Title',17),('Heading 1',13),('Heading 2',11.5),('Heading 3',11)]:
        if name in doc.styles:
            st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True
            st.paragraph_format.space_before=Pt(11);st.paragraph_format.space_after=Pt(5)
            st.paragraph_format.keep_with_next=True
    if 'Title' in doc.styles:
        doc.styles['Title'].paragraph_format.space_before=Pt(0)
    for name in ['Bibliography','Caption','Image Caption','Table Caption']:
        if name in doc.styles:
            st=doc.styles[name];st.font.size=Pt(9.5 if name=='Bibliography' else 10)
            st.paragraph_format.line_spacing=1.05;st.paragraph_format.space_after=Pt(5)
    for p in doc.paragraphs:
        p.paragraph_format.widow_control=True
        text=p.text.strip()
        if text.startswith('Table ') or text.startswith('Figure '):
            p.paragraph_format.keep_with_next=text.startswith('Table ')
            for r in p.runs:r.font.size=Pt(10)
        if p.style.name.startswith('Heading'):
            p.paragraph_format.space_before=Pt(12)
            p.paragraph_format.space_after=Pt(6)
        if p.style.name.startswith('Heading') or p.style.name=='Title':
            for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
            pr=p._p.find(qn('w:pPr'))
            if pr is not None:
                for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
        if p._p.xpath('.//w:drawing'):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_after=Pt(3)
    for ti,table in enumerate(doc.tables):
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.autofit=False
        widths={0:[1.5,2.75,2.55],1:[.85,1.15,1.85,1.45,1.5],2:[2.65,1.38,1.38,1.39],3:[1.72,2.87,2.21],4:[1.1,.65,1,1,1,1,1.05],5:[.9,1.05,1.05,1.25,1.25,1.3],6:[.9,.6,.8,1.2,.8,1.4,1.1],10:[1.0,1.0,.5,.65,.65,3.0],11:[2.6,.65,1.15,1.15,1.25]}.get(ti)
        if widths and len(widths)==len(table.columns):
            for column,width in zip(table.columns,widths):column.width=Inches(width)
            for row in table.rows:
                for cell,width in zip(row.cells,widths):cell.width=Inches(width)
        tblpr=table._tbl.tblPr
        borders=set_element(tblpr,'w:tblBorders',{})
        for edge in ['top','left','bottom','right','insideH','insideV']:
            set_element(borders,'w:'+edge,{'w:val':'single','w:sz':'4','w:color':'D9D9D9'})
        for idx,row in enumerate(table.rows):
            pr=row._tr.get_or_add_trPr()
            if idx==0:set_element(pr,'w:tblHeader',{'w:val':'true'})
            set_element(pr,'w:cantSplit',{})
            for cell in row.cells:
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr()
                margins=set_element(tcpr,'w:tcMar',{})
                for edge,val in [('top',80),('bottom',80),('left',95),('right',95)]:
                    set_element(margins,'w:'+edge,{'w:w':val,'w:type':'dxa'})
                set_element(tcpr,'w:shd',{'w:fill':'E8EDF2' if idx==0 else 'FFFFFF'})
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
                    p.paragraph_format.line_spacing=1.02
                    if len(table.rows)<=8:
                        # Keep short main tables with their explanatory note.
                        p.paragraph_format.keep_with_next=True
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if len(p.text)>24 else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name='Times New Roman';r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
                        if idx==0:r.font.bold=True
    doc.core_properties.title='Auditing financial data versions in bankruptcy prediction'
    doc.core_properties.subject='Anonymous research manuscript with reproducibility appendices'
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.save(args.output)
    xml=doc._element.xml
    report={'output':str(args.output),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),
            'inline_shapes':len(doc.inline_shapes),'native_equations':len(re.findall(r'<m:oMath[ >]',xml)),
            'word_count_markdown':len(re.findall(r"\b[\w'-]+\b",source)),
            'visual_qa':'PENDING_RENDER_AND_PAGE_INSPECTION'}
    (ROOT/'results/manuscript_build.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps(report))

if __name__=='__main__':main()
