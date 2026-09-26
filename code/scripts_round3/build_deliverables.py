"""Render the revised manuscript and supplement into Word and PDF deliverables.

The main text uses native Word equations and Times New Roman layout; the PDF is
produced through Word so that equations stay visible. Both outputs are written to
the delivery folder, and the resulting PDF text is checked afterwards by
pdf_qa.py.
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT.parent
SOURCE_ROUND = BASE / 'revision_20260922_round2'
MAN = ROOT / 'manuscript'
FIG_MAIN = SOURCE_ROUND / 'figures'
FIG_SUPP = SOURCE_ROUND / 'revision_20260922/figures'
TEMP = ROOT / 'temp'
DELIVERY = ROOT / 'delivery'
TEMP.mkdir(exist_ok=True)
DELIVERY.mkdir(exist_ok=True)

TITLE = ('Financial data version risk in bankruptcy model validation: '
         'separating scoring, development and evaluation effects')


def set_element(parent, tag, attrs):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    for key, value in attrs.items():
        el.set(qn(key), str(value))
    return el


def pandoc(source: Path, output: Path, resource_paths: list[Path]) -> None:
    exe = shutil.which('pandoc')
    if not exe:
        raise FileNotFoundError('pandoc is not available')
    command = [exe, str(source), '--standalone',
               '--resource-path=' + ';'.join(str(p) for p in resource_paths),
               '-M', 'lang=en-US', '-o', str(output)]
    run = subprocess.run(command, capture_output=True, text=True, encoding='utf-8', cwd=ROOT)
    (TEMP / f'pandoc_{source.stem}.log').write_text(run.stdout + '\n' + run.stderr,
                                                    encoding='utf-8')
    if run.returncode:
        raise RuntimeError(run.stderr)
    if 'Could not convert TeX math' in run.stderr:
        raise RuntimeError('native equation conversion failed: ' + run.stderr)


def style_document(doc: Document, widths: dict[int, list[float]]) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.text = ''
        run = footer.add_run()
        run.font.name = 'Times New Roman'
        run.font.size = Pt(9)
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        run._r.addnext(field)
    for style in doc.styles:
        if style.type in (1, 2):
            style.font.name = 'Times New Roman'
            style.font.color.rgb = RGBColor(0, 0, 0)
            rpr = style.element.find(qn('w:rPr'))
            if rpr is not None:
                color = rpr.find(qn('w:color'))
                if color is not None:
                    for attr in ['w:themeColor', 'w:themeTint', 'w:themeShade']:
                        color.attrib.pop(qn(attr), None)
    for name in ['Normal', 'Body Text', 'First Paragraph']:
        if name in doc.styles:
            st = doc.styles[name]
            st.font.size = Pt(11)
            st.paragraph_format.line_spacing = 1.15
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.widow_control = True
    for name, size in [('Title', 16), ('Heading 1', 12.5), ('Heading 2', 11.5), ('Heading 3', 11)]:
        if name in doc.styles:
            st = doc.styles[name]
            st.font.size = Pt(size)
            st.font.bold = True
            st.paragraph_format.space_before = Pt(12)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.keep_with_next = True
    if 'Title' in doc.styles:
        doc.styles['Title'].paragraph_format.space_before = Pt(0)
    for name in ['Caption', 'Image Caption', 'Table Caption']:
        if name in doc.styles:
            st = doc.styles[name]
            st.font.size = Pt(9.5)
            st.paragraph_format.line_spacing = 1.05
            st.paragraph_format.space_after = Pt(5)
    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
        text = p.text.strip()
        if text.startswith('Note:'):
            p.paragraph_format.keep_together = True
            p.paragraph_format.space_before = Pt(2)
            for r in p.runs:
                r.font.size = Pt(9)
        if re.match(r'^(Table|Figure) \d+\.', text):
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.size = Pt(9.5)
        if p._p.xpath('.//w:drawing'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(4)
        if p.style.name.startswith('Heading') or p.style.name == 'Title':
            for r in p.runs:
                r.font.color.rgb = RGBColor(0, 0, 0)
            pr = p._p.find(qn('w:pPr'))
            if pr is not None:
                for border in list(pr.findall(qn('w:pBdr'))):
                    pr.remove(border)
    for ti, table in enumerate(doc.tables):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths_i = widths.get(ti)
        if not widths_i:
            lengths = [max(len(row.cells[j].text) for row in table.rows)
                       for j in range(len(table.columns))]
            weights = [min(3.0, max(0.8, (v / 10) ** 0.55)) for v in lengths]
            widths_i = [6.6 * w / sum(weights) for w in weights]
        if len(widths_i) == len(table.columns):
            for column, width in zip(table.columns, widths_i):
                column.width = Inches(width)
            for row in table.rows:
                for cell, width in zip(row.cells, widths_i):
                    cell.width = Inches(width)
        tblpr = table._tbl.tblPr
        borders = set_element(tblpr, 'w:tblBorders', {})
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            set_element(borders, 'w:' + edge,
                        {'w:val': 'single', 'w:sz': '4', 'w:color': 'D9D9D9'})
        for idx, row in enumerate(table.rows):
            pr = row._tr.get_or_add_trPr()
            if idx == 0:
                set_element(pr, 'w:tblHeader', {'w:val': 'true'})
            set_element(pr, 'w:cantSplit', {})
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                margins = set_element(tcpr, 'w:tcMar', {})
                for edge, val in [('top', 70), ('bottom', 70), ('left', 90), ('right', 90)]:
                    set_element(margins, 'w:' + edge, {'w:w': val, 'w:type': 'dxa'})
                set_element(tcpr, 'w:shd',
                            {'w:fill': 'E8EDF2' if idx == 0 else 'FFFFFF'})
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.line_spacing = 1.02
                    p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if len(p.text) > 22
                                   else WD_ALIGN_PARAGRAPH.CENTER)
                    for r in p.runs:
                        r.font.name = 'Times New Roman'
                        r.font.size = Pt(8.5)
                        r.font.color.rgb = RGBColor(0, 0, 0)
                        if idx == 0:
                            r.font.bold = True


MAIN_WIDTHS = {
    0: [3.6, 1.0, 0.9, 1.1],
    1: [1.05] + [0.694] * 8,
    2: [1.15, 0.5, 1.35, 1.2, 1.35, 1.05],
    3: [1.07, 1.05, 1.02, 1.02, 1.42, 1.02],
    4: [1.05, 2.6, 1.05, 1.05, 0.85],
    5: [1.0, 1.85, 0.95, 0.75, 0.75, 0.75, 0.85],
}


def build_main() -> dict:
    source = MAN / 'manuscript.md'
    text = source.read_text(encoding='utf-8')
    if '{{' in text:
        raise SystemExit('unsubstituted placeholder remains in the manuscript')
    if re.search(r'\b(TODO|TBD|PLACEHOLDER|INSERT RESULT)\b', text, re.I):
        raise SystemExit('unresolved placeholder text')
    title_block = f'---\ntitle: "{TITLE}"\nlang: en-US\n---\n\n'
    staged = TEMP / 'manuscript_styled.md'
    staged.write_text(title_block + text, encoding='utf-8')
    intermediate = TEMP / 'manuscript_pandoc.docx'
    pandoc(staged, intermediate, [FIG_MAIN, MAN, ROOT])
    doc = Document(intermediate)
    style_document(doc, MAIN_WIDTHS)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = 'Anonymous research manuscript'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    out = DELIVERY / 'Financial_Data_Version_Risk_Manuscript.docx'
    doc.save(out)
    xml = doc._element.xml
    report = {'output': str(out), 'paragraphs': len(doc.paragraphs),
              'tables': len(doc.tables), 'figures': len(doc.inline_shapes),
              'native_equations': len(re.findall(r'<m:oMath[ >]', xml)),
              'words': len(re.findall(r"\b[\w']+\b", text))}
    return report


def build_supplement() -> dict:
    source = MAN / 'supplement.md'
    if not source.exists():
        return {'output': None, 'note': 'no supplement staged for this revision'}
    intermediate = TEMP / 'supplement_pandoc.docx'
    pandoc(source, intermediate, [FIG_SUPP, MAN, ROOT])
    doc = Document(intermediate)
    style_document(doc, {})
    doc.core_properties.title = 'Supplementary Methods and Results'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    out = DELIVERY / 'Financial_Data_Version_Risk_Supplement.docx'
    doc.save(out)
    return {'output': str(out), 'paragraphs': len(doc.paragraphs),
            'tables': len(doc.tables), 'figures': len(doc.inline_shapes)}


def to_pdf(docx: Path) -> str:
    """Convert a Word file to PDF through Word, returning the PDF path or a reason."""
    pdf = docx.with_suffix('.pdf')
    if pdf.exists():
        pdf.unlink()
    word = None
    try:
        import win32com.client  # type: ignore
        word = win32com.client.Dispatch('Word.Application')
    except Exception:
        try:
            import comtypes.client  # type: ignore
            word = comtypes.client.CreateObject('Word.Application')
        except Exception as exc:
            return f'unavailable: {exc}'
    try:
        word.Visible = False
        document = word.Documents.Open(str(docx), ReadOnly=True)
        document.ExportAsFixedFormat(str(pdf), 17)
        document.Close(False)
    except Exception as exc:
        return f'failed: {exc}'
    finally:
        try:
            word.Quit()
        except Exception:
            pass
    return str(pdf) if pdf.exists() else 'failed: no output'


if __name__ == '__main__':
    main_report = build_main()
    supp_report = build_supplement()
    main_pdf = to_pdf(Path(main_report['output']))
    supp_pdf = to_pdf(Path(supp_report['output'])) if supp_report.get('output') else None
    report = {'main': main_report, 'supplement': supp_report,
              'main_pdf': main_pdf, 'supplement_pdf': supp_pdf}
    (ROOT / 'results/build_report.json').write_text(json.dumps(report, indent=2),
                                                    encoding='utf-8')
    print(json.dumps(report, indent=2))
