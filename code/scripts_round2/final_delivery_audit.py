"""Bind the completed human-visible artifact review to frozen files.

This consolidates already performed computational and AI visual checks. It is
not human accounting certification, external empirical validation or training.
Package execution is separately certified after final ZIP creation.
"""
from pathlib import Path
import csv, hashlib, json, re, shutil, stat
from datetime import datetime, timezone
from docx import Document
from pypdf import PdfReader

R=Path(__file__).resolve().parents[1]
P=R.parent
M=R/'manuscript'

def sha(p):
    h=hashlib.sha256()
    with Path(p).open('rb') as f:
        for block in iter(lambda:f.read(1024*1024),b''):h.update(block)
    return h.hexdigest()

def load(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))

def page_record(folder,n,reviewer,basis):
    p=R/'qa'/folder/f'page-{n}.png'
    return dict(page=n,sha256=sha(p),reviewer=reviewer,basis=basis,status='PASS')

def main():
    protected={
        P/'datasets/model_cohort.parquet':'724122d3ca7a0a123de244112aa604d6da7a008497cc6ce526f8f07327916e21',
        P/'revision_20260922/manuscript/manuscript.md':'9c9b91d7a75d2fadea79fddebb00ec3b6b023c6883c4dba4e7e4664d7a8c17a8',
        P/'revision_20260922/manuscript/Auditing_Financial_Data_Versions_Revised.docx':'26d767f3aa67896836c455ace71922ca634fbcc039c28b3e8eacb7ca2f038ce0',
    }
    protection=[]
    for p,h in protected.items():
        actual=sha(p);assert actual==h,(p,actual)
        protection.append(dict(path=str(p.relative_to(P)),sha256=actual,status='UNCHANGED'))
    main_review=load(R/'results/source_semantics/visual_review.json')
    assert main_review['pdf_sha256']==sha(R/'qa/main_verified/Financial_Data_Version_Evaluation.pdf')
    for row in main_review['pages_inspected']:
        assert row['sha256']==sha(R/'qa/main_verified'/f"page-{row['page']}.png")
    main_pages=[]
    for n in range(1,10):
        if n<=6:
            assert sha(R/'qa/main_verified'/f'page-{n}.png')==sha(R/'qa/main_final'/f'page-{n}.png')
            basis='Actual visual inspection of main_final plus byte-identical final PNG'
        else:basis='Actual visual inspection of main_verified PNG'
        main_pages.append(page_record('main_verified',n,'root AI reviewer',basis))
    main_pages += [dict(row,reviewer='finance_route AI reviewer',status='PASS') for row in main_review['pages_inspected']]
    supp_pages=[page_record('supplement_verified',n,'root AI reviewer','Actual final PNG visual inspection') for n in range(1,8)]
    for a,b,who in [(8,16,'technology_route AI reviewer'),(17,26,'ai_route AI reviewer')]:
        supp_pages += [page_record('supplement_verified',n,who,'Actual final PNG visual inspection; corresponding reviewer report retained') for n in range(a,b+1)]
    review_files=[R/'results/supplement_tables/visual_qa_final_8_16.json',R/'reproducibility/supplement_visual_qa_final_17_26.json']
    for p in review_files:
        assert p.exists()
        assert sha(R/'qa/supplement_verified/Supplementary_Methods_and_Results.pdf') in p.read_text(encoding='utf-8')
    numerical_files=[
        'results/development_controls/numerical_audit.json',
        'results/development_controls/fit_date_independent_audit.json',
        'results/evaluation_diagnostics/independent_metric_audit.json',
        'results/scale_only/independent_audit.json',
    ]
    numerical=[]
    for s in numerical_files:
        d=load(R/s);assert d['status']=='PASS'
        numerical.append(dict(path=s,sha256=sha(R/s),status='PASS',details=d))
    outputs=[]
    for stem,source,folder,table_n,fig_n,pages,equations in [
        ('Financial_Data_Version_Evaluation','manuscript','main_verified',6,4,main_pages,25),
        ('Supplementary_Methods_and_Results','supplement','supplement_verified',25,0,supp_pages,26),
    ]:
        dx=M/(stem+'.docx');pdf_qa=R/'qa'/folder/(stem+'.pdf');pdf_out=M/(stem+'.pdf')
        shutil.copy2(pdf_qa,pdf_out)
        doc=Document(dx)
        assert len(doc.tables)==table_n and len(doc.inline_shapes)==fig_n
        assert len(re.findall(r'<m:oMath[ >]',doc._element.xml))==equations
        pdf=PdfReader(pdf_out).pages;assert len(pdf)==len(pages)
        assert all(len(p.extract_text().strip())>20 for p in pdf)
        body='\n'.join(p.text for p in doc.paragraphs)
        assert not re.search(r'\b(?:TODO|TBD|PLACEHOLDER)\b',body)
        assert not re.search(r'[CDE]:[\\/]',body)
        md=(M/(source+'.md')).read_text(encoding='utf-8')
        cited=set(re.findall(r'@([A-Za-z][\w:.-]*)',md))
        bib=(M/'references.bib').read_text(encoding='utf-8')
        keys=set(re.findall(r'@\w+\s*\{\s*([^,]+),',bib))
        assert cited<=keys,(cited-keys)
        assert len(cited)==(40 if source=='manuscript' else 9)
        build=load(R/'results'/(source+'_build.json'))
        build.update(visual_qa='PASS_ALL_RENDERED_PAGES_INSPECTED',pdf_pages=len(pdf),final_pdf_sha256=sha(pdf_out),final_docx_sha256=sha(dx))
        (R/'results'/(source+'_build.json')).write_text(json.dumps(build,indent=2),encoding='utf-8')
        outputs.append(dict(stem=stem,docx_sha256=sha(dx),pdf_sha256=sha(pdf_out),markdown_sha256=sha(M/(source+'.md')),tables=table_n,figures=fig_n,native_equations=equations,cited_sources=len(cited),pdf_pages=len(pdf),page_review=pages))
    pending=R/'results/source_semantics/author_pending_all_consequential_sources.csv'
    rows=list(csv.DictReader(pending.open(encoding='utf-8-sig',newline='')))
    assert len(rows)==59
    # All confirmation columns intentionally stay empty; never synthesize signoff.
    signature_cols=['author_name','author_review_date','author_selected_decision','author_source_locator','author_explanation','author_signature']
    assert signature_cols
    assert all(not row[c].strip() for row in rows for c in signature_cols)
    raw=Path('<archived-record-root>/...')
    raw_files=[p for p in raw.rglob('*') if p.is_file()]
    writable=[str(p) for p in raw_files if p.stat().st_mode & stat.S_IWRITE]
    assert not writable,writable
    report=dict(
        status='PASS_ARTIFACTS_AND_COMPUTATIONAL_REVISION',
        generated_utc=datetime.now(timezone.utc).isoformat(),
        submission_decision='CONDITIONAL',
        submission_stage='research revision; not final journal submission authorization',
        scope='AI-assisted computation/source assessment and visual artifact review; no human accounting certification',
        protected_inputs=protection,artifacts=outputs,numerical_audits=numerical,
        supplementary_reviewer_reports=[dict(path=str(p.relative_to(R)),sha256=sha(p)) for p in review_files],
        source_confirmation=dict(rows=59,completed_human_signoffs=0,sha256=sha(pending)),
        raw_storage=dict(files=len(raw_files),all_read_only=True,location=str(raw)),
        package_validation='Reported separately in reproducibility/portable_validation.json after final ZIP build; no package hash self-reference',
        remaining=['Human accounting/XBRL source review','Independent empirical prediction or complete alternate-extraction replication','Public repository deposition and actual journal-specific submission checks'],
    )
    (R/'results/final_revision_audit.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
    (R/'results/final_revision_audit.md').write_text(
        '# Final revision artifact audit\n\n'
        'PASS for the completed computational revision and artifact review. Submission remains CONDITIONAL.\n\n'
        'Main manuscript: 18 pages, 6 tables, 4 figures, 25 native Office equation objects, 40 cited sources. '
        'Supplement: 26 pages, 25 tables, 26 native equation objects, 9 cited sources. All 44 rendered pages were actually inspected by AI reviewers; unchanged main pages were bound by byte identity. '
        'No clipping or unresolved mathematical rendering issue remains.\n\n'
        'The original cohort and prior main Markdown/Word match their pre-revision hashes. Numerical check reports pass. '
        'The 59-row source-confirmation form has no completed human signatures; original issuer-instance authentication and new empirical validation are not certified. '
        'Final package execution and ZIP identity are recorded separately after packaging.\n',encoding='utf-8')
    print(json.dumps(dict(status=report['status'],pages=44,source_rows=59,raw_files=len(raw_files)),ensure_ascii=False))

if __name__=='__main__':main()
